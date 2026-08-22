"""Uploaded files -> normalized text -> ~1500-char segments.

The text produced here is what source spans point at later, so `normalize`
never lowercases or reflows: it only fixes line endings and whitespace.
Segments are numbered per document here; global `seq` is assigned when they
are written to the DB.
"""
import io
import re
from dataclasses import dataclass
from pathlib import PurePosixPath

from app.config import get_settings


class IngestError(ValueError):
    """User-facing: the message is safe to show in the UI."""


@dataclass
class IngestedDoc:
    filename: str
    content_type: str
    text: str
    char_count: int
    segments: list[str]


_CONTENT_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_TEXT_EXTS = (".txt", ".md")
_BINARY_CONTROL_RATIO = 0.02


def _ext(filename: str) -> str:
    return PurePosixPath(filename).suffix.lower()


def content_type_for(filename: str) -> str:
    return _CONTENT_TYPES.get(_ext(filename), "application/octet-stream")


# --- extraction ------------------------------------------------------------


def _decode_text(data: bytes) -> str:
    text = None
    for enc in ("utf-8", "utf-8-sig"):
        try:
            text = data.decode(enc, errors="strict")
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("latin-1")
    if text:
        controls = sum(1 for c in text if ord(c) < 32 and c not in "\t\r\n")
        if controls / len(text) > _BINARY_CONTROL_RATIO:
            raise IngestError("Looks like a binary file")
    return text


def _pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise IngestError("Not a valid PDF file")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append(notes)
        if parts:
            slides.append("\n".join(parts))
    return "\n\n".join(slides)


def extract_text(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if ext in _TEXT_EXTS:
        text = _decode_text(data)
    elif ext in (".pdf", ".docx", ".pptx"):
        reader = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx}[ext]
        try:
            text = reader(data)
        except IngestError:
            raise
        except Exception as e:  # corrupt/odd files from the wild
            raise IngestError(f"Could not read {filename}: {type(e).__name__}") from e
    else:
        raise IngestError(
            f"Unsupported file type '{ext or '(none)'}' for {filename}; "
            "use .txt, .md, .pdf, .docx or .pptx"
        )
    if not text.strip():
        raise IngestError(f"No readable text found in {filename}")
    return text


# --- normalize / segment ---------------------------------------------------

_ZERO_WIDTH = re.compile("[​‌‍﻿]")
_SPACE_RUN = re.compile(r"[ \t]+")
_TRAILING_SPACE = re.compile(r"[ \t]+$", re.MULTILINE)
_BLANK_RUN = re.compile(r"\n{3,}")
_PARA_BREAK = re.compile(r"\n\s*\n")
# Split after sentence punctuation, keeping the separator so pieces re-join exactly.
_SENTENCE_SEP = re.compile(r"((?<=[.!?])\s+)")


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _ZERO_WIDTH.sub("", text)
    text = _SPACE_RUN.sub(" ", text)
    text = _TRAILING_SPACE.sub("", text)
    text = _BLANK_RUN.sub("\n\n", text)
    return text.strip()


def _split_long_paragraph(paragraph: str, hard_max: int) -> list[str]:
    """Greedy sentence packing; a sentence longer than hard_max is hard-cut."""
    tokens = _SENTENCE_SEP.split(paragraph)  # [sentence, sep, sentence, sep, ...]
    sentences = tokens[0::2]
    seps = tokens[1::2] + [""]
    out: list[str] = []
    cur = ""
    pending_sep = ""
    for sentence, sep in zip(sentences, seps):
        if len(sentence) > hard_max:
            if cur:
                out.append(cur)
                cur = ""
            out.extend(sentence[i : i + hard_max] for i in range(0, len(sentence), hard_max))
            pending_sep = sep
            continue
        candidate = cur + pending_sep + sentence if cur else sentence
        if len(candidate) > hard_max:
            out.append(cur)
            cur = sentence
        else:
            cur = candidate
        pending_sep = sep
    if cur:
        out.append(cur)
    return [p for p in out if p.strip()]


def segment(text: str, target: int = 1500, hard_max: int = 2400) -> list[str]:
    paragraphs = [p.strip() for p in _PARA_BREAK.split(text) if p.strip()]
    pieces: list[str] = []
    for p in paragraphs:
        if len(p) > hard_max:
            pieces.extend(_split_long_paragraph(p, hard_max))
        else:
            pieces.append(p)

    segments: list[str] = []
    cur: list[str] = []
    cur_len = 0
    for piece in pieces:
        added = len(piece) + (2 if cur else 0)
        if cur and cur_len + added > target:
            segments.append("\n\n".join(cur))
            cur, cur_len = [], 0
            added = len(piece)
        cur.append(piece)
        cur_len += added
    if cur:
        segments.append("\n\n".join(cur))
    return segments


# --- entry point -----------------------------------------------------------


def _make_doc(filename: str, text: str, target: int) -> IngestedDoc:
    return IngestedDoc(
        filename=filename,
        content_type=content_type_for(filename),
        text=text,
        char_count=len(text),
        segments=segment(text, target=target),
    )


def ingest_files(
    items: list[tuple[str, bytes]],
    pasted_text: str | None,
    *,
    min_total_chars: int,
    max_total_chars: int,
    segment_target: int | None = None,
) -> list[IngestedDoc]:
    target = segment_target or get_settings().segment_chars
    docs: list[IngestedDoc] = []
    for filename, data in items:
        text = normalize(extract_text(filename, data))
        docs.append(_make_doc(filename, text, target))
    if pasted_text and pasted_text.strip():
        docs.append(_make_doc("pasted.txt", normalize(pasted_text), target))

    total = sum(d.char_count for d in docs)
    if total < min_total_chars:
        raise IngestError(
            f"Not enough text: {total:,} characters, need at least {min_total_chars:,}. "
            "Add more material."
        )
    if total > max_total_chars:
        raise IngestError(
            f"Too much text: {total:,} characters, the limit is {max_total_chars:,}. "
            "Remove some files or split the course."
        )
    return docs
