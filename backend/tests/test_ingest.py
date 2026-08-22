"""Ingest tests. Office files are built in memory with the same libraries the
service reads them with, so nothing on disk is needed.
"""
import io

import docx
import pytest
from pptx import Presentation
from pypdf import PdfWriter

from app.services import ingest


def _docx_bytes() -> bytes:
    d = docx.Document()
    d.add_paragraph("Hello from Word.")
    d.add_paragraph("")
    d.add_paragraph("Second paragraph here.")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "cell one"
    table.cell(0, 1).text = "cell two"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = "Body bullet"
    slide.notes_slide.notes_text_frame.text = "Speaker notes here"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Second slide"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _blank_pdf_bytes() -> bytes:
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


# --- extract_text ----------------------------------------------------------


def test_txt_utf8():
    assert ingest.extract_text("notes.txt", "héllo wörld".encode()) == "héllo wörld"


def test_md_utf8_with_bom():
    out = ingest.extract_text("notes.MD", "﻿hello".encode())
    assert out.lstrip("﻿") == "hello"


def test_txt_latin1_fallback():
    assert ingest.extract_text("notes.txt", "café".encode("latin-1")) == "café"


def test_txt_binary_sniff_rejects():
    with pytest.raises(ingest.IngestError, match="binary"):
        ingest.extract_text("weird.txt", bytes(range(256)))


def test_unsupported_extension():
    with pytest.raises(ingest.IngestError, match="Unsupported"):
        ingest.extract_text("archive.zip", b"PK\x03\x04")


def test_empty_txt_is_error():
    with pytest.raises(ingest.IngestError, match="No readable text"):
        ingest.extract_text("empty.txt", b"   \n  ")


def test_docx_paragraphs_and_tables():
    out = ingest.extract_text("lecture.docx", _docx_bytes())
    assert "Hello from Word." in out
    assert "Second paragraph here." in out
    assert "cell one" in out and "cell two" in out
    assert "\n\n\n" not in out  # empty paragraph skipped


def test_pptx_shapes_and_notes():
    out = ingest.extract_text("deck.pptx", _pptx_bytes())
    assert "Slide title" in out
    assert "Body bullet" in out
    assert "Speaker notes here" in out
    assert "Second slide" in out
    assert out.index("Slide title") < out.index("Second slide")


def test_pdf_blank_page_has_no_text():
    data = _blank_pdf_bytes()
    assert data.startswith(b"%PDF-")
    with pytest.raises(ingest.IngestError, match="No readable text"):
        ingest.extract_text("blank.pdf", data)


def test_pdf_requires_magic():
    with pytest.raises(ingest.IngestError):
        ingest.extract_text("fake.pdf", b"this is not a pdf")


def test_corrupt_docx_is_user_error():
    with pytest.raises(ingest.IngestError):
        ingest.extract_text("broken.docx", b"definitely not a zip")


def test_content_type_for():
    assert ingest.content_type_for("a.txt") == "text/plain"
    assert ingest.content_type_for("a.PDF") == "application/pdf"
    assert ingest.content_type_for("a.bin") == "application/octet-stream"


# --- normalize -------------------------------------------------------------


def test_normalize_line_endings():
    assert ingest.normalize("a\r\nb\rc") == "a\nb\nc"


def test_normalize_zero_width_chars():
    assert ingest.normalize("a​b‌c‍d﻿e") == "abcde"


def test_normalize_collapses_spaces_and_trailing():
    assert ingest.normalize("a  \t b   \nc \n") == "a b\nc"


def test_normalize_collapses_blank_lines():
    assert ingest.normalize("a\n\n\n\n\nb\n\n\nc") == "a\n\nb\n\nc"


def test_normalize_never_lowercases_or_reflows():
    text = "Hello World.\nSecond LINE stays."
    assert ingest.normalize(text) == text


# --- segment ---------------------------------------------------------------


def test_segment_short_text_is_one_segment():
    assert ingest.segment("Just a short note.") == ["Just a short note."]


def test_segment_empty_text_is_empty():
    assert ingest.segment("   ") == []


def test_segment_keeps_paragraphs_whole():
    paragraphs = [f"P{i} " + ("word " * 79).strip() for i in range(10)]  # ~400 chars each
    assert all(390 <= len(p) <= 410 for p in paragraphs)
    segs = ingest.segment("\n\n".join(paragraphs))
    assert len(segs) >= 2
    assert all(len(s) <= 2400 for s in segs)
    for seg in segs:
        for part in seg.split("\n\n"):
            assert part in paragraphs
    assert "\n\n".join(segs) == "\n\n".join(paragraphs)


def test_segment_splits_long_paragraph_on_sentences():
    sentences = [f"Sentence number {i} says something mildly interesting." for i in range(120)]
    text = " ".join(sentences)
    assert len(text) > 6000
    segs = ingest.segment(text)
    assert len(segs) >= 3
    for seg in segs:
        assert len(seg) <= 2400
        assert seg.rstrip()[-1] in ".!?"
    joined = " ".join(segs)
    for s in sentences:
        assert s in joined


def test_segment_hard_cuts_a_giant_sentence():
    text = "x" * 5000
    segs = ingest.segment(text)
    assert all(len(s) <= 2400 for s in segs)
    assert "".join(segs) == text


# --- ingest_files ----------------------------------------------------------


def test_ingest_files_pasted_text_becomes_doc():
    docs = ingest.ingest_files([], "Some pasted text. " * 20, min_total_chars=10, max_total_chars=10_000)
    assert len(docs) == 1
    d = docs[0]
    assert d.filename == "pasted.txt"
    assert d.content_type == "text/plain"
    assert d.char_count == len(d.text)
    assert d.segments == [d.text]


def test_ingest_files_mixes_files_and_paste():
    docs = ingest.ingest_files(
        [("a.txt", b"File text here. " * 10), ("b.docx", _docx_bytes())],
        "pasted",
        min_total_chars=10,
        max_total_chars=10_000,
    )
    assert [d.filename for d in docs] == ["a.txt", "b.docx", "pasted.txt"]
    assert all(d.segments for d in docs)


def test_ingest_files_min_total():
    with pytest.raises(ingest.IngestError, match="at least"):
        ingest.ingest_files([("a.txt", b"tiny")], None, min_total_chars=100, max_total_chars=1000)


def test_ingest_files_max_total():
    with pytest.raises(ingest.IngestError, match="limit"):
        ingest.ingest_files([("a.txt", b"x" * 500)], None, min_total_chars=1, max_total_chars=100)


def test_ingest_files_nothing_is_min_error():
    with pytest.raises(ingest.IngestError):
        ingest.ingest_files([], "", min_total_chars=1, max_total_chars=100)
